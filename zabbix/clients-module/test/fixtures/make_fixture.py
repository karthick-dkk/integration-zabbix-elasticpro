"""
The workbook as another tool sees it, and two workbooks another tool wrote.

    php test/xlsx.test.php writes template.xlsx and export.xlsx here when EP_FIXTURE_OUT=test/fixtures
    docker run --rm -v "$PWD":/w -w /w/test/fixtures python:3.12-slim sh -c "pip -q install openpyxl && python make_fixture.py"

(a) opens our template.xlsx and export.xlsx with openpyxl and checks the sheets, headers, data
    validations (type, formula, error style), notes, frozen pane and the collapsed capacity group;
(b) fills template.xlsx as a person would in Excel and saves openpyxl-filled.xlsx;
(c) writes openpyxl-mixed.xlsx from nothing: a date cell, numbers, a boolean, empty cells, an extra
    attribute column, a lower-case sheet name and a sheet that is not read.
test/xlsx.test.php then reads (b) and (c) with our reader and checks what the import would get.

libreoffice-export.xlsx is export.xlsx opened and saved again by LibreOffice (heavy, so by hand):
    docker run --rm -v "$PWD":/w -w /w debian:bookworm-slim sh -c "apt-get update -qq && \
      apt-get install -y -qq --no-install-recommends libreoffice-calc && \
      soffice --headless --convert-to xlsx --outdir lo export.xlsx && cp lo/export.xlsx libreoffice-export.xlsx"
"""
import datetime
import sys

import openpyxl
from openpyxl.utils import get_column_letter

failed = []
passed = []


def check(what, ok, detail=None):
    passed.append(what) if ok else None
    if not ok:
        failed.append(what)
        print("FAIL:", what, "" if detail is None else f"— {detail}")


def headers(ws):
    return [c.value for c in ws[1]]


for name in ("template.xlsx", "export.xlsx"):
    wb = openpyxl.load_workbook(name)
    check(f"{name}: three sheets in order", wb.sheetnames == ["Clients", "Servers", "Sample"], wb.sheetnames)
    cl, sv, sa = wb["Clients"], wb["Servers"], wb["Sample"]
    h = headers(cl)
    check(f"{name}: Clients headers start as designed", h[:4] == ["Client *", "Type ▾", "Purchased by ▾", "Purchased"], h[:4])
    check(f"{name}: Contract ends and Agent port are there", "Contract ends" in h and "Agent port" in h)
    check(f"{name}: a capacity column per role part", "ES Data Hot · Servers" in h and "ES Data Hot · Other disks" in h and "AIML · Memory GB" in h)
    check(f"{name}: freeze panes", cl.freeze_panes == "B2" and sv.freeze_panes == "A2", (cl.freeze_panes, sv.freeze_panes))
    for col in range(1, len(h) + 1):
        if h[col - 1] is None:
            continue
        cell = cl.cell(row=1, column=col)
        if str(h[col - 1]).startswith("Choices: "):
            continue
        check(f"{name}: Clients {h[col - 1]} has a note", cell.comment is not None and len(cell.comment.text) > 10)
    for col in range(1, len(headers(sv)) + 1):
        check(f"{name}: Servers {headers(sv)[col - 1]} has a note", sv.cell(row=1, column=col).comment is not None)
    # The capacity group: outline level 1, hidden; the base columns are not.
    hot = h.index("ES Data Hot · Servers") + 1
    dim = cl.column_dimensions[get_column_letter(hot)]
    check(f"{name}: capacity columns are outline level 1 and collapsed", dim.outline_level == 1 and dim.hidden, (dim.outline_level, dim.hidden))
    check(f"{name}: ES URL is not in the group", cl.column_dimensions[get_column_letter(h.index("ES URL *") + 1)].outline_level in (0, None))
    choices = [i + 1 for i, v in enumerate(h) if v and str(v).startswith("Choices: ")]
    check(f"{name}: a hidden Choices column for Monitored by", any(h[i - 1] == "Choices: Monitored by" for i in choices))
    for i in choices:
        d = cl.column_dimensions[get_column_letter(i)]
        check(f"{name}: {h[i - 1]} is hidden and outside the group", d.hidden and not d.outline_level)
    dvs = {}
    for dv in cl.data_validations.dataValidation:
        dvs[str(dv.sqref)] = dv
    def dv_for(ws, header):
        col = get_column_letter(headers(ws).index(header) + 1)
        for dv in ws.data_validations.dataValidation:
            if str(dv.sqref) == f"{col}2:{col}2000":
                return dv
        return None
    t = dv_for(cl, "Type ▾")
    check(f"{name}: Type is a list of the client types", t is not None and t.type == "list" and t.formula1 == '"CI,DI,On-Prem"' and t.errorStyle == "stop"
          and t.showErrorMessage and "CI, DI or On-Prem" in (t.error or ""), t and (t.formula1, t.errorStyle, t.error))
    m = dv_for(cl, "Monitored by ▾")
    check(f"{name}: Monitored by points at its hidden column", m is not None and m.formula1.startswith("Clients!$") and m.type == "list", m and m.formula1)
    for hd, f in (("Alert DL ▾", '"Yes,No"'), ("Weekly report ▾", '"Yes,No"'), ("Purchased by ▾", '"Storage,Devices"'),
                  ("Password in ▾", '"Vault,Zabbix secret"'), ("Jump TLS ▾", '"Yes,No"'), ("S3 access ▾", '"Zabbix server role,Access key"')):
        d = dv_for(cl, hd)
        check(f"{name}: {hd} list", d is not None and d.formula1 == f and d.errorStyle == "stop" and d.prompt, d and d.formula1)
    c = dv_for(sv, "Client ▾ *")
    check(f"{name}: Servers Client reads the Clients sheet", c is not None and c.formula1 == "Clients!$A$2:$A$2000", c and c.formula1)
    for n in (1, 2, 3, 4):
        hd = f"Role {n} ▾" + (" *" if n == 1 else "")
        d = dv_for(sv, hd)
        check(f"{name}: {hd} is a list of the roles", d is not None and d.type == "list" and ("ES Data Hot" in d.formula1 or d.formula1.startswith("Clients!$")), d and d.formula1)
    check(f"{name}: Sample banner", str(sa["A1"].value).startswith("SAMPLE — not imported") and "A1:K1" in [str(r) for r in sa.merged_cells.ranges]
          and sa["A1"].fill.fgColor.rgb in ("FFFFF4D6", "00FFF4D6"), sa["A1"].fill.fgColor.rgb)
    check(f"{name}: Sample has notes on the tricky cells", all(sa[ref].comment is not None for ref in ("D4", "K4", "C10", "D10")))

ex = openpyxl.load_workbook("export.xlsx")["Clients"]
eh = headers(ex)
check("export: Contract ends is a real date", isinstance(ex.cell(row=2, column=eh.index("Contract ends") + 1).value, datetime.datetime),
      ex.cell(row=2, column=eh.index("Contract ends") + 1).value)
check("export: Purchased is a number", isinstance(ex.cell(row=2, column=4).value, (int, float)), ex.cell(row=2, column=4).value)
check("export: a formula-like lead is a string, not a formula", ex.cell(row=2, column=eh.index("SOC / Manager lead") + 1).data_type == "s"
      and ex.cell(row=2, column=eh.index("SOC / Manager lead") + 1).value == "=cmd|' /C calc'!A0")

# (b) filled in, as a person would.
wb = openpyxl.load_workbook("template.xlsx")
cl, sv = wb["Clients"], wb["Servers"]
h = headers(cl)
col = {v: i + 1 for i, v in enumerate(h) if v}
rows = [
    {"Client *": "Bajaj", "Type ▾": "DI", "Purchased by ▾": "Storage", "Purchased": 20000, "SOC / Manager lead": "Asha Rao",
     "Cluster DL": "soc-bajaj@example.com", "Alert DL ▾": "Yes", "Weekly report ▾": "Yes", "Contract ends": datetime.datetime(2027, 3, 31),
     "ES URL *": "https://es.bajaj.example:9200", "ES user": "elastic", "Password in ▾": "Vault", "Vault path": "secret/elasticpro/bajaj:password",
     "Monitored by ▾": "Zabbix server", "S3 bucket": "bajaj-archive", "S3 region": "ap-south-1", "S3 access ▾": "Zabbix server role",
     "ES Data Hot · Servers": 2, "ES Data Hot · CPU cores": 16, "ES Data Hot · Memory GB": 64, "ES Data Hot · Disk / GB": 200,
     "ES Data Hot · Other disks": "/data=4000; /logs=500"},
    {"Client *": "Wayne", "Type ▾": "On-Prem", "Purchased by ▾": "Devices", "Purchased": 300, "Cluster DL": "soc-wayne@example.com",
     "Alert DL ▾": "No", "Weekly report ▾": "No", "ES URL *": "https://es.wayne.example:9200", "Password in ▾": "Zabbix secret",
     "Monitored by ▾": "proxy-mumbai-01", "Agent port": 10051},
]
for r, row in enumerate(rows, start=2):
    for k, v in row.items():
        cl.cell(row=r, column=col[k], value=v)
cl.cell(row=2, column=col["Contract ends"]).number_format = "yyyy-mm-dd"
sh = headers(sv)
sv.cell(row=1, column=len(sh) + 1, value="Rack")
sc = {v: i + 1 for i, v in enumerate(headers(sv)) if v}
servers = [
    {"Client ▾ *": "Bajaj", "IP *": "10.0.4.11", "Role 1 ▾ *": "ES Data Hot", "Rack": "R4"},
    {"Client ▾ *": "Bajaj", "IP *": "10.0.4.12", "Role 1 ▾ *": "ES Master", "Role 2 ▾": "ES Coordination", "Rack": "R4"},
    {"Client ▾ *": "Wayne", "IP *": "10.0.5.1", "Role 1 ▾ *": "ES Data", "Role 2 ▾": "Parser", "Role 3 ▾": "Forwarder", "Role 4 ▾": "Engine",
     "Services": "nginx; kafka", "Notes": "single node"},
]
for r, row in enumerate(servers, start=2):
    for k, v in row.items():
        sv.cell(row=r, column=sc[k], value=v)
wb.save("openpyxl-filled.xlsx")

# (c) from nothing.
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Clients"
ws.append(["Client", "Type", "Contract ends", "Purchased by", "Purchased", "Alert DL", "Cluster DL", "Agent port", "ES URL", "Favourite colour"])
ws.append(["Mixed", "on prem", datetime.date(2026, 12, 31), "storage", 1500.0, True, "dl@example.com", 10050, None, "blue"])
ws.append([None, None, None, None, None, None, None, None, None, None])
ws.append(["Mixed2", "CI", "2027-01-15", None, 12.5, False, None, None, "https://es.mixed2.example:9200", None])
ws["C2"].number_format = "dd/mm/yyyy"
srv = wb.create_sheet("servers")
srv.append(["Client", "IP", "Role 1", "Role 2", "Owner team", "Services"])
srv.append(["Mixed", "10.30.0.1", "es_data", "parser", "team-a", None])
srv.append(["Mixed", "10.30.0.2", "Forwarder", None, None, "nginx"])
wb.create_sheet("Notes").append(["not read", "at all"])
wb.save("openpyxl-mixed.xlsx")

print(f"{len(passed)} passed, {len(failed)} failed")
sys.exit(1 if failed else 0)
